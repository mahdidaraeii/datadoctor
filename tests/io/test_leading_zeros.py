import json

import openpyxl
import pandas as pd

from datadoctor import load_dataset
from datadoctor.io import leading_zeros

N = 60


def planted_csv(tmp_path):
    """60 rows with a column for each case. What is recorded, and why, is listed in the test."""
    # sku is first: it is text, so the columns after it sit at a different place among the numeric
    # columns than among all of them.
    lines = ["sku,zip,zip_gaps,mixed_width,age,balance,ratio,flag,active"]
    for k in range(N):
        zip_code = ["01234", "56789", "00042", "90210"][k % 4]  # 5 wide, no gaps, half with zeros
        zip_gaps = ["01111", "22222", "", "03333"][k % 4]  # 5 wide with blanks, half with zeros
        mixed_width = ["007", "12", "0042", "8"][k % 4]  # zeros, but no single width
        balance = ["-5", "007"][k % 2]  # a leading zero, but next to negatives: not codes
        ratio = ["05.5", "0.25"][k % 2]  # a leading zero, but fractions: not codes
        flag = ["0", "1"][k % 2]  # a lone zero is not a leading zero
        sku = "?" if k % 30 == 0 else f"{k:03d}"  # zero padded, but read as text because of "?"
        active = ["true", "false"][k % 2]
        lines.append(
            f"{sku},{zip_code},{zip_gaps},{mixed_width},{20 + k},{balance},{ratio},{flag},{active}"
        )
    path = tmp_path / "planted.csv"
    path.write_text("\n".join(lines) + "\n", encoding="utf-8", newline="")
    return path


class TestCsv:
    def test_columns_that_lost_leading_zeros_are_recorded_and_no_other(self, tmp_path):
        dataset = load_dataset(planted_csv(tmp_path))

        # Each column cycles through four values, so half of the 60 rows have a leading zero.
        assert dataset.provenance.leading_zeros == {
            "zip": {"values": 30, "checked": N, "width": 5},
            "zip_gaps": {"values": 30, "checked": N, "width": 5},
            "mixed_width": {"values": 30, "checked": N},
        }

    def test_the_record_describes_what_already_happened_to_the_loaded_data(self, tmp_path):
        dataset = load_dataset(planted_csv(tmp_path))

        assert pd.api.types.is_integer_dtype(dataset.data["zip"])
        assert dataset.data["zip"][0] == 1234
        assert dataset.data["sku"][1] == "001"  # a text column keeps its zeros, so is not recorded

    def test_only_the_first_100000_rows_are_checked_and_the_record_says_how_many(self, tmp_path):
        rows = [f"{k + 1:06d},{'000007' if k >= 100_000 else '7'}" for k in range(100_010)]
        path = tmp_path / "long.csv"
        path.write_text("early,late\n" + "\n".join(rows) + "\n", encoding="utf-8", newline="")

        recorded = load_dataset(path).provenance.leading_zeros

        # "early" counts 1 to 100,010 written 6 wide. Of the first 100,000 rows, 99,999 have a
        # leading zero, and 100,000 does not. "late" has zeros only after those rows, so it is
        # not seen.
        assert recorded == {"early": {"values": 99_999, "checked": 100_000, "width": 6}}

    def test_a_file_without_candidate_columns_records_nothing_and_is_read_only_once(
        self, tmp_path, monkeypatch
    ):
        path = tmp_path / "text.csv"
        path.write_text("a,b\nx,0.5\ny,0.25\n", encoding="utf-8", newline="")
        reads = []
        read_csv = pd.read_csv
        monkeypatch.setattr(pd, "read_csv", lambda *a, **k: reads.append(1) or read_csv(*a, **k))

        assert load_dataset(path).provenance.leading_zeros == {}
        assert len(reads) == 1


class TestExcel:
    def test_zeros_in_text_cells_are_found_and_display_formats_are_not_data(self, tmp_path):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.append(["as_text", "formatted", "mixed"])
        for k, code in enumerate(["01234", "09317", "12695", "02713"]):
            sheet.append([code, int(code), code if k % 2 else int(code)])
        for row in sheet.iter_rows(min_row=2, min_col=2, max_col=2):
            for cell in row:
                cell.number_format = "00000"  # shows 01234, and stores the number 1234
        path = tmp_path / "zips.xlsx"
        book.save(path)

        dataset = load_dataset(path)

        assert pd.api.types.is_integer_dtype(dataset.data["as_text"])
        assert dataset.provenance.leading_zeros == {
            "as_text": {"values": 3, "checked": 4, "width": 5},
            "mixed": {"values": 2, "checked": 4},
        }

    def test_only_the_first_rows_are_checked_here_too(self, tmp_path, monkeypatch):
        monkeypatch.setattr(leading_zeros, "CHECK_ROWS", 3)
        book = openpyxl.Workbook()
        book.active.append(["early", "late"])
        for k in range(6):
            book.active.append(["01234", "01234" if k >= 3 else "1234"])
        path = tmp_path / "long.xlsx"
        book.save(path)

        recorded = load_dataset(path).provenance.leading_zeros

        assert recorded == {"early": {"values": 3, "checked": 3, "width": 5}}


class TestFormatsThatKeepText:
    def test_json_keeps_zero_padded_strings_so_nothing_is_lost_or_recorded(self, tmp_path):
        path = tmp_path / "zips.json"
        path.write_text(json.dumps([{"zip": "01234"}, {"zip": "09317"}]), encoding="utf-8")

        dataset = load_dataset(path)

        assert dataset.data["zip"].tolist() == ["01234", "09317"]
        assert dataset.provenance.leading_zeros == {}

    def test_parquet_keeps_zero_padded_strings_so_nothing_is_lost_or_recorded(self, tmp_path):
        path = tmp_path / "zips.parquet"
        pd.DataFrame({"zip": ["01234", "09317"]}).to_parquet(path)

        dataset = load_dataset(path)

        assert dataset.data["zip"].tolist() == ["01234", "09317"]
        assert dataset.provenance.leading_zeros == {}
